import logging
import os

logger = logging.getLogger(__name__)

# Files Gemini can read natively via the Files API (images, pdf)
GEMINI_NATIVE_EXTENSIONS = {".pdf", ".png", ".jpg", ".jpeg", ".webp", ".gif"}
DOCX_EXTENSIONS = {".docx"}
PLAIN_TEXT_EXTENSIONS = {".txt", ".csv", ".md", ".json"}
XLSX_EXTENSIONS = {".xlsx", ".xlsm"}

def classify_file(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext in GEMINI_NATIVE_EXTENSIONS:
        return "gemini_native"
    if ext in DOCX_EXTENSIONS:
        return "docx"
    if ext in XLSX_EXTENSIONS:
        return "xlsx"
    if ext in PLAIN_TEXT_EXTENSIONS:
        return "plain_text"
    return "unsupported"


def extract_text_from_docx(path: str) -> str:
    try:
        import docx
    except ImportError:
        logger.warning("python-docx not installed; cannot parse .docx files. `pip install python-docx`.")
        return ""
    try:
        doc = docx.Document(path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
        return "\n".join(parts)
    except Exception as e:
        logger.error(f"Failed to parse docx {path}: {e}")
        return ""


def extract_text_from_plain(path: str) -> str:
    try:
        with open(path, "r", errors="ignore") as f:
            return f.read()
    except Exception as e:
        logger.error(f"Failed to read text file {path}: {e}")
        return ""


def extract_text_from_xlsx(path: str) -> str:
    try:
        import openpyxl
    except ImportError:
        logger.warning("openpyxl not installed; cannot parse .xlsx files. `pip install openpyxl`.")
        return ""
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        parts = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None and str(c).strip()]
                if cells:
                    parts.append(" ".join(cells))
        return "\n".join(parts)
    except Exception as e:
        logger.error(f"Failed to parse xlsx {path}: {e}")
        return ""